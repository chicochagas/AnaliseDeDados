"""
Corrige as inconsistências da planilha de vendas e gera uma planilha tratada.

Entrada : data/planilha_vendas.csv
Saída   : data/planilha_vendas_tratada.xlsx
          - Vendas Tratadas : dados corrigidos + coluna de alertas
          - Resumo          : indicadores calculados por fórmula
          - Log de Correções: cada alteração feita (valor original -> corrigido)
          - Pendências      : itens que exigem confirmação humana

Uso: python scripts/corrigir_vendas.py [entrada.csv] [saida.xlsx]

Também pode ser importado (usado pelo app Streamlit):
    brutos = ler_csv(conteudo_bytes)
    resultado = limpar(brutos)
    gerar_excel(resultado, destino)
"""

import csv
import io
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

RAIZ = Path(__file__).resolve().parent.parent
ENTRADA_PADRAO = RAIZ / "data" / "planilha_vendas.csv"
SAIDA_PADRAO = RAIZ / "data" / "planilha_vendas_tratada.xlsx"

# ---------------------------------------------------------------- domínios
COLUNAS_OBRIGATORIAS = ["ID Pedido", "Nome do Cliente", "Região", "Representante", "Data do Pedido",
                        "Receita", "Produto", "Status", "Observações"]
REGIOES = ["Norte", "Nordeste", "Centro-Oeste", "Sudeste", "Sul"]
STATUS_VALIDOS = {
    "fechado ganho": "Fechado Ganho",
    "closed won": "Fechado Ganho",
    "ganho": "Fechado Ganho",
    "fechado perdido": "Fechado Perdido",
    "closed lost": "Fechado Perdido",
    "perdido": "Fechado Perdido",
}
MESES = {
    "jan": 1, "janeiro": 1, "fev": 2, "fevereiro": 2, "mar": 3, "março": 3, "marco": 3,
    "abr": 4, "abril": 4, "mai": 5, "maio": 5, "jun": 6, "junho": 6,
    "jul": 7, "julho": 7, "ago": 8, "agosto": 8, "set": 9, "setembro": 9,
    "out": 10, "outubro": 10, "nov": 11, "novembro": 11, "dez": 12, "dezembro": 12,
}
# Marcações de qualidade que não devem ficar em "Observações"
MARCADORES_QUALIDADE = {"duplicado", "possível duplicado", "receita ausente", "sem representante"}
PARTICULAS = {"da", "de", "do", "das", "dos", "e"}


class ErroEntrada(ValueError):
    """Arquivo enviado não tem o formato esperado."""


def sem_acento(txt):
    return "".join(c for c in unicodedata.normalize("NFD", txt) if unicodedata.category(c) != "Mn")


def moeda_br(valor):
    return "R$ " + f"{valor:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")


def chave_id(pid):
    """Ordena IDs numéricos antes de IDs em texto."""
    return (0, pid, "") if isinstance(pid, int) else (1, 0, str(pid))


# ---------------------------------------------------------------- normalizadores
def normalizar_nome(valor):
    partes = valor.strip().lower().split()
    return " ".join(p if (p in PARTICULAS and i > 0) else p.capitalize() for i, p in enumerate(partes))


def normalizar_regiao(valor):
    chave = sem_acento(valor.strip().lower()).replace(" ", "-")
    for r in REGIOES:
        if sem_acento(r.lower()) == chave:
            return r
    return ""


def normalizar_status(valor):
    return STATUS_VALIDOS.get(valor.strip().lower(), "")


def converter_data(valor):
    v = valor.strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(v, fmt)
        except ValueError:
            pass
    m = re.fullmatch(r"(\d{1,2})\s+([a-zà-ú]+)\.?\s+(\d{4})", v.lower())
    if m and m.group(2) in MESES:
        try:
            return datetime(int(m.group(3)), MESES[m.group(2)], int(m.group(1)))
        except ValueError:
            return None
    return None


def converter_receita(valor):
    """Retorna float, None se vazio; levanta ValueError se não for número."""
    v = valor.replace("R$", "").strip()
    if not v:
        return None
    if "," in v:  # formato brasileiro 1.234,56
        v = v.replace(".", "").replace(",", ".")
    return float(v)


# ---------------------------------------------------------------- leitura
def ler_csv(origem):
    """Lê o CSV a partir de um caminho ou de bytes (upload). Detecta codificação e separador."""
    dados = Path(origem).read_bytes() if isinstance(origem, (str, Path)) else bytes(origem)
    for codificacao in ("utf-8-sig", "cp1252"):
        try:
            texto = dados.decode(codificacao)
            break
        except UnicodeDecodeError:
            continue
    if not texto.strip():
        raise ErroEntrada("O arquivo está vazio.")

    primeira_linha = texto.splitlines()[0]
    separador = max([",", ";", "\t"], key=primeira_linha.count)
    leitor = csv.DictReader(io.StringIO(texto, newline=""), delimiter=separador)
    leitor.fieldnames = [c.strip() for c in (leitor.fieldnames or [])]

    faltando = [c for c in COLUNAS_OBRIGATORIAS if c not in leitor.fieldnames]
    if faltando:
        raise ErroEntrada("Colunas obrigatórias ausentes: " + ", ".join(faltando)
                          + ". Colunas encontradas: " + ", ".join(leitor.fieldnames))
    return [{c: (linha.get(c) or "") for c in COLUNAS_OBRIGATORIAS}
            for linha in leitor if any((v or "").strip() for v in linha.values() if isinstance(v, str))]


# ---------------------------------------------------------------- limpeza
def limpar(brutos):
    """
    Aplica todas as correções. Retorna um dicionário com:
      registros  : linhas corrigidas (dicts)
      log        : alterações feitas [{id, coluna, original, corrigido, regra, categoria}]
      pendencias : itens para revisão humana [{ids, tipo, descricao, acao}]
      alertas    : {id: [mensagens]}
      removidos  : IDs removidos por duplicidade exata
      lidas      : quantidade de linhas lidas
    """
    log, pendencias, alertas = [], [], defaultdict(list)

    def registrar(pid, coluna, original, corrigido, regra, categoria):
        if str(original) != str(corrigido):
            log.append({"id": pid, "coluna": coluna, "original": original, "corrigido": corrigido,
                        "regra": regra, "categoria": categoria})

    def pendente(pid, tipo, descricao, acao, alerta=None):
        pendencias.append({"ids": str(pid), "tipo": tipo, "descricao": descricao, "acao": acao})
        if alerta:
            alertas[pid].append(alerta)

    # ---- 1. padronização campo a campo
    registros = []
    contagem_ids = Counter(b["ID Pedido"].strip() for b in brutos)
    for n, b in enumerate(brutos, start=2):
        id_txt = b["ID Pedido"].strip()
        pid = int(id_txt) if id_txt.isdigit() else (id_txt or f"(linha {n})")
        if not id_txt.isdigit():
            pendente(pid, "ID inválido", f"ID '{id_txt}' não é numérico", "Corrigir o ID na origem", "ID inválido")
        elif contagem_ids[id_txt] > 1:
            pendente(pid, "ID repetido", f"ID {id_txt} aparece {contagem_ids[id_txt]} vezes", "Verificar lançamentos",
                     "ID repetido")
        r = {"ID Pedido": pid, "_original": b}

        r["Nome do Cliente"] = normalizar_nome(b["Nome do Cliente"])
        registrar(pid, "Nome do Cliente", b["Nome do Cliente"], r["Nome do Cliente"],
                  "Capitalização padronizada", "Capitalização fora do padrão")

        r["Região"] = normalizar_regiao(b["Região"])
        if r["Região"]:  # região inválida fica vazia e é tratada na etapa de inferência
            registrar(pid, "Região", b["Região"], r["Região"], "Grafia padronizada para lista oficial",
                      "Região fora do padrão")

        r["Representante"] = normalizar_nome(b["Representante"])
        registrar(pid, "Representante", b["Representante"], r["Representante"],
                  "Capitalização padronizada", "Capitalização fora do padrão")

        data = converter_data(b["Data do Pedido"])
        r["Data do Pedido"] = data
        if data is None:
            pendente(pid, "Data inválida", f"Data '{b['Data do Pedido']}' não reconhecida", "Corrigir manualmente",
                     "Data inválida")
        else:
            registrar(pid, "Data do Pedido", b["Data do Pedido"], data.strftime("%d/%m/%Y"),
                      "Data convertida para DD/MM/AAAA", "Data em formato não padrão")

        try:
            r["Receita"] = converter_receita(b["Receita"])
        except ValueError:
            r["Receita"] = None
            pendente(pid, "Receita inválida", f"Receita '{b['Receita']}' não é um número",
                     "Corrigir o valor na origem", "Receita inválida")
        if r["Receita"] is not None:
            registrar(pid, "Receita", b["Receita"], r["Receita"], "Texto convertido em número", "Receita como texto")

        r["Produto"] = b["Produto"].strip()

        obs = b["Observações"].strip()
        r["Status"] = normalizar_status(b["Status"])
        if b["Status"].strip() and not r["Status"]:
            pendente(pid, "Status desconhecido", f"Status '{b['Status']}' fora da lista", "Classificar manualmente",
                     "Status desconhecido")
        elif not r["Status"] and normalizar_status(obs):
            # status digitado na coluna errada
            r["Status"] = normalizar_status(obs)
            registrar(pid, "Status", b["Status"] or "(vazio)", r["Status"],
                      "Status recuperado da coluna Observações", "Status na coluna errada")
            obs = ""
        elif b["Status"].strip().lower() in ("closed won", "closed lost"):
            registrar(pid, "Status", b["Status"], r["Status"], "Traduzido para português", "Status em outro idioma")
        else:
            registrar(pid, "Status", b["Status"], r["Status"], "Capitalização padronizada",
                      "Capitalização fora do padrão")

        if obs.lower() in MARCADORES_QUALIDADE:
            registrar(pid, "Observações", obs, "(vazio)", "Marcação de qualidade removida (tratada em Alertas/Log)",
                      "Marcador de erro em Observações")
            obs = ""
        r["Observações"] = obs
        registros.append(r)

    # ---- 2. duplicados
    campos = ["Nome do Cliente", "Região", "Representante", "Data do Pedido", "Receita", "Produto", "Status"]
    removidos, vistos = set(), {}
    for r in registros:
        assinatura = tuple(str(r[c]).lower() for c in campos)
        if r["Data do Pedido"] is not None and assinatura in vistos:  # linha idêntica = duplicado exato
            removidos.add(r["ID Pedido"])
            registrar(r["ID Pedido"], "(linha inteira)", "registro", "REMOVIDO",
                      f"Duplicado exato do pedido {vistos[assinatura]['ID Pedido']}", "Duplicado exato")
        else:
            vistos.setdefault(assinatura, r)
    registros = [r for r in registros if r["ID Pedido"] not in removidos]

    grupos = defaultdict(list)
    for r in registros:
        if r["Receita"] is not None and r["Nome do Cliente"]:
            grupos[(r["Nome do Cliente"].lower(), r["Produto"].lower(), r["Receita"])].append(r)
    for itens in grupos.values():
        if len(itens) < 2:
            continue
        itens.sort(key=lambda x: chave_id(x["ID Pedido"]))
        ids = ", ".join(str(x["ID Pedido"]) for x in itens)
        status = {x["Status"] for x in itens if x["Status"]}
        extra = " (status divergentes: " + " x ".join(sorted(status)) + ")" if len(status) > 1 else ""
        for x in itens:
            alertas[x["ID Pedido"]].append(f"Possível duplicado: pedidos {ids}{extra}")
        pendente(ids, "Possível duplicado",
                 f"{itens[0]['Nome do Cliente']} / {itens[0]['Produto']} / {moeda_br(itens[0]['Receita'])} "
                 f"em datas diferentes{extra}",
                 "Confirmar com o comercial se é recompra ou lançamento repetido")

    # ---- 3. campos vazios (inferência)
    regioes_por_rep = defaultdict(Counter)
    reps_por_regiao = defaultdict(Counter)
    for r in registros:
        if r["Representante"] and r["Região"]:
            regioes_por_rep[r["Representante"]][r["Região"]] += 1
            reps_por_regiao[r["Região"]][r["Representante"]] += 1

    for r in registros:
        pid = r["ID Pedido"]
        if not r["Região"]:
            opcoes = regioes_por_rep.get(r["Representante"], {})
            if len(opcoes) == 1:
                r["Região"] = next(iter(opcoes))
                registrar(pid, "Região", r["_original"]["Região"] or "(vazio)", r["Região"],
                          f"Inferida: {r['Representante']} atua apenas em {r['Região']}", "Campo vazio preenchido")
                alertas[pid].append("Região inferida pelo representante")
            elif r["_original"]["Região"].strip():
                pendente(pid, "Região inválida", f"Região '{r['_original']['Região']}' fora da lista oficial",
                         "Corrigir manualmente", "Região inválida")
            else:
                pendente(pid, "Região ausente", "Não foi possível inferir", "Preencher manualmente", "Região ausente")
        if not r["Representante"]:
            opcoes = reps_por_regiao.get(r["Região"], {})
            if len(opcoes) == 1:
                r["Representante"] = next(iter(opcoes))
                registrar(pid, "Representante", "(vazio)", r["Representante"],
                          f"Inferido: único representante da região {r['Região']}", "Campo vazio preenchido")
                alertas[pid].append("Representante inferido pela região")
            else:
                pendente(pid, "Representante ausente", "Não foi possível inferir", "Preencher manualmente",
                         "Representante ausente")
        if r["Receita"] is None and not r["_original"]["Receita"].strip():
            pendente(pid, "Receita ausente", f"Pedido '{r['Status'] or 'sem status'}' sem valor "
                                             f"({r['Nome do Cliente']})",
                     "Buscar valor no contrato/nota fiscal", "Receita ausente")
        if not r["Status"] and not r["_original"]["Status"].strip():
            pendente(pid, "Status ausente", "Sem status nem indicação em Observações", "Preencher manualmente",
                     "Status ausente")

    return {"registros": registros, "log": log, "pendencias": pendencias, "alertas": dict(alertas),
            "removidos": removidos, "lidas": len(brutos)}


# ---------------------------------------------------------------- planilha
FONTE = "Arial"
CAB = Font(name=FONTE, bold=True, color="FFFFFF")
CAB_FILL = PatternFill("solid", start_color="1F4E78")
NORMAL = Font(name=FONTE)
NEGRITO = Font(name=FONTE, bold=True)
ALERTA_FILL = PatternFill("solid", start_color="FFF2CC")
BORDA = Border(bottom=Side(style="thin", color="BFBFBF"))
MOEDA = '"R$" #,##0.00;-"R$" #,##0.00;"-"'


def _cabecalho(ws, colunas, larguras):
    ws.append(colunas)
    for i, (c, w) in enumerate(zip(ws[1], larguras), 1):
        c.font, c.fill = CAB, CAB_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"


def gerar_excel(resultado, destino):
    """Grava a planilha tratada em `destino` (caminho ou objeto tipo arquivo, ex.: BytesIO)."""
    registros, alertas = resultado["registros"], resultado["alertas"]
    wb = Workbook()
    wb.calculation.fullCalcOnLoad = True  # Excel recalcula as fórmulas ao abrir

    # --- Vendas Tratadas
    ws = wb.active
    ws.title = "Vendas Tratadas"
    colunas = ["ID Pedido", "Nome do Cliente", "Região", "Representante", "Data do Pedido",
               "Receita (R$)", "Produto", "Status", "Observações", "Alertas"]
    _cabecalho(ws, colunas, [11, 22, 14, 18, 14, 15, 22, 17, 20, 55])
    for r in registros:
        ws.append([r["ID Pedido"], r["Nome do Cliente"], r["Região"] or None, r["Representante"] or None,
                   r["Data do Pedido"], r["Receita"], r["Produto"], r["Status"] or None, r["Observações"] or None,
                   "; ".join(alertas.get(r["ID Pedido"], [])) or None])
    for linha in ws.iter_rows(min_row=2):
        tem_alerta = linha[9].value is not None
        for c in linha:
            c.font, c.border = NORMAL, BORDA
            if tem_alerta:
                c.fill = ALERTA_FILL
        linha[4].number_format = "DD/MM/YYYY"
        linha[5].number_format = MOEDA
    ultima = max(ws.max_row, 2)
    ws.auto_filter.ref = f"A1:J{ultima}"

    # --- Resumo (fórmulas sobre a aba de dados)
    rs = wb.create_sheet("Resumo")
    rs.column_dimensions["A"].width = 34
    rs.column_dimensions["B"].width = 18
    rs.column_dimensions["C"].width = 18
    V = "'Vendas Tratadas'!"
    F_REC, F_STA, F_REG = f"{V}$F$2:$F${ultima}", f"{V}$H$2:$H${ultima}", f"{V}$C$2:$C${ultima}"

    rs["A1"] = "Resumo da Planilha de Vendas Tratada"
    rs["A1"].font = Font(name=FONTE, bold=True, size=14)
    indicadores = [
        ("Total de pedidos", f"=COUNTA({V}$A$2:$A${ultima})", "0"),
        ("Receita total (todos os status)", f"=SUM({F_REC})", MOEDA),
        ("Receita Fechado Ganho", f'=SUMIFS({F_REC},{F_STA},"Fechado Ganho")', MOEDA),
        ("Pedidos Fechado Ganho", f'=COUNTIFS({F_STA},"Fechado Ganho")', "0"),
        ("Pedidos Fechado Perdido", f'=COUNTIFS({F_STA},"Fechado Perdido")', "0"),
        ("Taxa de conversão", "=IFERROR(B6/(B6+B7),0)", "0.0%"),
        ("Ticket médio (ganhos com receita)",
         f'=IFERROR(B5/COUNTIFS({F_STA},"Fechado Ganho",{F_REC},"<>"),0)', MOEDA),
        ("Pedidos sem receita", f'=COUNTBLANK({F_REC})', "0"),
        ("Pedidos com alerta", f'=COUNTA({V}$J$2:$J${ultima})', "0"),
    ]
    for i, (rot, form, fmt) in enumerate(indicadores, start=3):
        rs.cell(i, 1, rot).font = NORMAL
        c = rs.cell(i, 2, form)
        c.font, c.number_format = NEGRITO, fmt

    base = 3 + len(indicadores) + 1
    for j, t in enumerate(["Região", "Receita Ganha", "Pedidos"], 1):
        c = rs.cell(base, j, t)
        c.font, c.fill = CAB, CAB_FILL
    for k, reg in enumerate(REGIOES, start=base + 1):
        rs.cell(k, 1, reg).font = NORMAL
        c = rs.cell(k, 2, f'=SUMIFS({F_REC},{F_REG},A{k},{F_STA},"Fechado Ganho")')
        c.font, c.number_format = NORMAL, MOEDA
        rs.cell(k, 3, f"=COUNTIFS({F_REG},A{k})").font = NORMAL
    tot = base + len(REGIOES) + 1
    rs.cell(tot, 1, "Total").font = NEGRITO
    c = rs.cell(tot, 2, f"=SUM(B{base + 1}:B{tot - 1})")
    c.font, c.number_format = NEGRITO, MOEDA
    rs.cell(tot, 3, f"=SUM(C{base + 1}:C{tot - 1})").font = NEGRITO
    rs.cell(tot + 2, 1, "Obs.: possíveis duplicados foram mantidos e estão sinalizados; "
                        "confirme-os na aba Pendências antes de usar os totais.").font = \
        Font(name=FONTE, italic=True, color="7F7F7F")

    # --- Log de Correções
    lg = wb.create_sheet("Log de Correções")
    _cabecalho(lg, ["ID Pedido", "Coluna", "Valor Original", "Valor Corrigido", "Regra Aplicada", "Categoria"],
               [11, 18, 24, 24, 60, 30])
    for item in sorted(resultado["log"], key=lambda x: (chave_id(x["id"]), x["coluna"])):
        lg.append([item["id"], item["coluna"], str(item["original"]), str(item["corrigido"]), item["regra"],
                   item["categoria"]])
    for linha in lg.iter_rows(min_row=2):
        for c in linha:
            c.font = NORMAL
    lg.auto_filter.ref = f"A1:F{max(lg.max_row, 2)}"

    # --- Pendências
    pd_ = wb.create_sheet("Pendências")
    _cabecalho(pd_, ["ID(s) Pedido", "Tipo", "Descrição", "Ação Sugerida", "Resolvido? (S/N)"], [16, 20, 70, 50, 16])
    for p in resultado["pendencias"]:
        pd_.append([p["ids"], p["tipo"], p["descricao"], p["acao"], None])
    for linha in pd_.iter_rows(min_row=2):
        for c in linha:
            c.font = NORMAL
        linha[4].fill = PatternFill("solid", start_color="FFFF00")
    pd_.cell(pd_.max_row + 2, 1, "Preencha a coluna amarela após validar cada item com a área comercial.").font = \
        Font(name=FONTE, italic=True, color="7F7F7F")

    wb.save(destino)


# ---------------------------------------------------------------- linha de comando
def main(argv):
    entrada = Path(argv[0]) if len(argv) > 0 else ENTRADA_PADRAO
    saida = Path(argv[1]) if len(argv) > 1 else SAIDA_PADRAO
    brutos = ler_csv(entrada)
    resultado = limpar(brutos)
    saida.parent.mkdir(parents=True, exist_ok=True)
    gerar_excel(resultado, saida)
    print(f"Linhas lidas: {resultado['lidas']} | removidas: {len(resultado['removidos'])} "
          f"| finais: {len(resultado['registros'])}")
    print(f"Correções registradas: {len(resultado['log'])} | pendências: {len(resultado['pendencias'])}")
    print(f"Arquivo gerado: {saida}")


if __name__ == "__main__":
    main(sys.argv[1:])
