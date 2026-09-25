"""
Testes de validação do script scripts/corrigir_vendas.py.

O script é executado uma vez numa pasta temporária (o arquivo entregue em data/
não é sobrescrito) e a planilha gerada é conferida contra o CSV original e
contra cada item do diagnóstico.

Uso: python -m pytest tests -v
"""

import csv
import runpy
import sys
from collections import defaultdict
from datetime import datetime
from pathlib import Path

import pytest
from openpyxl import load_workbook

RAIZ = Path(__file__).resolve().parent.parent
SCRIPT = RAIZ / "scripts" / "corrigir_vendas.py"
ENTRADA = RAIZ / "data" / "planilha_vendas.csv"
ENTREGUE = RAIZ / "data" / "planilha_vendas_tratada.xlsx"

REGIOES = {"Norte", "Nordeste", "Centro-Oeste", "Sudeste", "Sul"}
STATUS = {"Fechado Ganho", "Fechado Perdido"}
PARTICULAS = {"da", "de", "do", "das", "dos", "e"}
MARCADORES = {"duplicado", "possível duplicado", "receita ausente", "sem representante",
              "fechado ganho", "fechado perdido", "closed won"}
COLUNAS = ["ID Pedido", "Nome do Cliente", "Região", "Representante", "Data do Pedido",
           "Receita (R$)", "Produto", "Status", "Observações", "Alertas"]


# ---------------------------------------------------------------- fixtures
@pytest.fixture(scope="session")
def execucao(tmp_path_factory):
    saida = tmp_path_factory.mktemp("saida") / "tratada.xlsx"
    argv = sys.argv
    sys.argv = [str(SCRIPT), str(ENTRADA), str(saida)]
    try:
        ns = runpy.run_path(str(SCRIPT), run_name="__main__")
    finally:
        sys.argv = argv
    return ns, load_workbook(saida)


@pytest.fixture(scope="session")
def ns(execucao):
    return execucao[0]


@pytest.fixture(scope="session")
def wb(execucao):
    return execucao[1]


@pytest.fixture(scope="session")
def originais():
    with open(ENTRADA, encoding="utf-8", newline="") as f:
        return {int(r["ID Pedido"]): r for r in csv.DictReader(f)}


@pytest.fixture(scope="session")
def vendas(wb):
    ws = wb["Vendas Tratadas"]
    return [dict(zip(COLUNAS, linha)) for linha in ws.iter_rows(min_row=2, values_only=True)]


@pytest.fixture(scope="session")
def por_id(vendas):
    return {v["ID Pedido"]: v for v in vendas}


@pytest.fixture(scope="session")
def log(wb):
    return list(wb["Log de Correções"].iter_rows(min_row=2, values_only=True))


@pytest.fixture(scope="session")
def pendencias(wb):
    return [l for l in wb["Pendências"].iter_rows(min_row=2, values_only=True) if l[1]]


def alertas(venda):
    return venda["Alertas"] or ""


# ---------------------------------------------------------------- funções de normalização
@pytest.mark.parametrize("entrada, esperado", [
    ("15/01/2024", datetime(2024, 1, 15)),
    ("2024-01-22", datetime(2024, 1, 22)),
    ("5 jan. 2024", datetime(2024, 1, 5)),
    ("3 março 2024", datetime(2024, 3, 3)),
    ("2 maio 2024", datetime(2024, 5, 2)),
    ("data errada", None),
])
def test_converter_data(ns, entrada, esperado):
    assert ns["converter_data"](entrada) == esperado


@pytest.mark.parametrize("entrada, esperado", [
    ("R$ 4320.00", 4320.0),
    ("R$ 1.234,56", 1234.56),
    ("990", 990.0),
    ("", None),
])
def test_converter_receita(ns, entrada, esperado):
    assert ns["converter_receita"](entrada) == esperado


@pytest.mark.parametrize("entrada, esperado", [
    ("SUL", "Sul"), ("sudeste", "Sudeste"), ("Centro-oeste", "Centro-Oeste"),
    ("centro oeste", "Centro-Oeste"), ("", ""), ("Marte", ""),
])
def test_normalizar_regiao(ns, entrada, esperado):
    assert ns["normalizar_regiao"](entrada) == esperado


@pytest.mark.parametrize("entrada, esperado", [
    ("fechado ganho", "Fechado Ganho"), ("closed won", "Fechado Ganho"),
    ("Fechado Perdido", "Fechado Perdido"), ("", ""), ("em negociação", ""),
])
def test_normalizar_status(ns, entrada, esperado):
    assert ns["normalizar_status"](entrada) == esperado


@pytest.mark.parametrize("entrada, esperado", [
    ("PEDRO OLIVEIRA", "Pedro Oliveira"), ("maria souza", "Maria Souza"),
    ("DÉRICO MENDES", "Dérico Mendes"), ("ana DA silva", "Ana da Silva"),
])
def test_normalizar_nome(ns, entrada, esperado):
    assert ns["normalizar_nome"](entrada) == esperado


# ---------------------------------------------------------------- estrutura da planilha
def test_abas_existem(wb):
    assert wb.sheetnames == ["Vendas Tratadas", "Resumo", "Log de Correções", "Pendências"]


def test_cabecalho(wb):
    assert [c.value for c in wb["Vendas Tratadas"][1]] == COLUNAS


def test_nenhum_pedido_perdido(vendas, originais, log):
    """Todo ID original está na planilha ou foi removido com registro no log."""
    saida = [v["ID Pedido"] for v in vendas]
    removidos = {l[0] for l in log if l[3] == "REMOVIDO"}
    assert len(saida) == len(set(saida)), "IDs repetidos na saída"
    assert set(saida) | removidos == set(originais)
    assert not set(saida) & removidos


# ---------------------------------------------------------------- itens do diagnóstico
def test_duplicado_exato_removido(por_id, log):
    assert 1005 not in por_id
    assert 1001 in por_id
    assert any(l[0] == 1005 and "1001" in l[4] for l in log if l[3] == "REMOVIDO")


def test_sem_duplicados_exatos(vendas):
    chaves = [(v["Nome do Cliente"].lower(), v["Produto"], v["Receita (R$)"], v["Data do Pedido"]) for v in vendas]
    assert len(chaves) == len(set(chaves))


def test_possiveis_duplicados_sinalizados(vendas, pendencias):
    grupos = defaultdict(list)
    for v in vendas:
        grupos[(v["Nome do Cliente"].lower(), v["Produto"], v["Receita (R$)"])].append(v["ID Pedido"])
    suspeitos = [sorted(ids) for ids in grupos.values() if len(ids) > 1]
    assert sorted(suspeitos) == [[1004, 1017], [1006, 1029], [1007, 1025], [1008, 1014]]
    ids_pendentes = {p[0] for p in pendencias if p[1] == "Possível duplicado"}
    for ids in suspeitos:
        for pid in ids:
            assert "Possível duplicado" in alertas(next(v for v in vendas if v["ID Pedido"] == pid))
        assert ", ".join(map(str, ids)) in ids_pendentes


def test_status_divergente_sinalizado(por_id):
    assert "status divergentes" in alertas(por_id[1007])
    assert "status divergentes" in alertas(por_id[1025])


@pytest.mark.parametrize("coluna", ["Nome do Cliente", "Representante"])
def test_capitalizacao_padronizada(vendas, coluna):
    for v in vendas:
        for i, palavra in enumerate(v[coluna].split()):
            if i > 0 and palavra in PARTICULAS:
                continue
            assert palavra[0].isupper() and palavra[1:] == palavra[1:].lower(), \
                f"{v['ID Pedido']}: '{v[coluna]}' fora do padrão"


def test_regiao_valida_e_preenchida(vendas):
    for v in vendas:
        assert v["Região"] in REGIOES, f"{v['ID Pedido']}: região '{v['Região']}'"


def test_representante_preenchido(vendas):
    assert all(v["Representante"] for v in vendas)


def test_status_valido_e_preenchido(vendas):
    for v in vendas:
        assert v["Status"] in STATUS, f"{v['ID Pedido']}: status '{v['Status']}'"


def test_closed_won_traduzido(por_id):
    assert por_id[1019]["Status"] == "Fechado Ganho"


def test_status_recuperado_de_observacoes(por_id, log):
    assert por_id[1007]["Status"] == "Fechado Perdido"
    assert por_id[1007]["Observações"] is None
    assert any(l[0] == 1007 and l[1] == "Status" and "Observações" in l[4] for l in log)


@pytest.mark.parametrize("pid, coluna, valor, alerta", [
    (1006, "Região", "Centro-Oeste", "Região inferida"),
    (1016, "Região", "Sul", "Região inferida"),
    (1027, "Representante", "Sara Jorge", "Representante inferido"),
])
def test_campos_inferidos(por_id, log, pid, coluna, valor, alerta):
    assert por_id[pid][coluna] == valor
    assert alerta in alertas(por_id[pid])
    assert any(l[0] == pid and l[1] == coluna and l[4].startswith("Inferid") for l in log)


def test_datas_convertidas(wb, vendas):
    for v in vendas:
        assert isinstance(v["Data do Pedido"], datetime), f"{v['ID Pedido']}: data não convertida"
        assert datetime(2024, 1, 1) <= v["Data do Pedido"] <= datetime(2024, 12, 31)
    for c in wb["Vendas Tratadas"]["E"][1:]:
        assert c.number_format == "DD/MM/YYYY"


def test_datas_iguais_ao_original(por_id, originais, ns):
    for pid, v in por_id.items():
        assert v["Data do Pedido"] == ns["converter_data"](originais[pid]["Data do Pedido"])


def test_receita_numerica(wb, vendas):
    for v in vendas:
        assert v["Receita (R$)"] is None or isinstance(v["Receita (R$)"], (int, float))
    for c in wb["Vendas Tratadas"]["F"][1:]:
        assert "R$" in c.number_format


def test_receita_preservada(por_id, originais):
    for pid, v in por_id.items():
        bruto = originais[pid]["Receita"].replace("R$", "").strip()
        assert v["Receita (R$)"] == (float(bruto) if bruto else None)


def test_receita_ausente_sinalizada(vendas, pendencias):
    sem_receita = {v["ID Pedido"] for v in vendas if v["Receita (R$)"] is None}
    assert sem_receita == {1009, 1021}
    for v in vendas:
        if v["ID Pedido"] in sem_receita:
            assert "Receita ausente" in alertas(v)
    assert {p[0] for p in pendencias if p[1] == "Receita ausente"} == {"1009", "1021"}


def test_observacoes_sem_marcadores(vendas):
    for v in vendas:
        obs = (v["Observações"] or "").lower()
        assert obs not in MARCADORES, f"{v['ID Pedido']}: marcador '{obs}' em Observações"


def test_observacao_legitima_mantida(por_id):
    assert por_id[1003]["Observações"] == "cliente recorrente"


# ---------------------------------------------------------------- rastreabilidade
@pytest.mark.parametrize("coluna_saida, coluna_csv", [
    ("Nome do Cliente", "Nome do Cliente"), ("Região", "Região"),
    ("Representante", "Representante"), ("Status", "Status"),
])
def test_toda_alteracao_esta_no_log(por_id, originais, log, coluna_saida, coluna_csv):
    registradas = {(l[0], l[1]) for l in log}
    for pid, v in por_id.items():
        if originais[pid][coluna_csv].strip() != v[coluna_saida]:
            assert (pid, coluna_csv) in registradas, f"{pid}/{coluna_csv} alterado sem log"


def test_linhas_com_alerta_destacadas(wb):
    for linha in wb["Vendas Tratadas"].iter_rows(min_row=2):
        cor = linha[0].fill.start_color.rgb
        if linha[9].value:
            assert cor.endswith("FFF2CC"), f"{linha[0].value}: alerta sem destaque"
        else:
            assert not cor.endswith("FFF2CC")


# ---------------------------------------------------------------- resumo
def test_resumo_usa_formulas(wb):
    rs = wb["Resumo"]
    for linha in range(3, 12):
        assert str(rs.cell(linha, 2).value).startswith("="), f"Resumo!B{linha} não é fórmula"


def test_resumo_cobre_todas_as_linhas(wb):
    ultima = wb["Vendas Tratadas"].max_row
    for linha in range(3, 12):
        formula = str(wb["Resumo"].cell(linha, 2).value)
        if "Vendas Tratadas" in formula:
            assert f"${ultima}" in formula, f"Resumo!B{linha} não vai até a linha {ultima}"


def test_totais_esperados(vendas):
    ganho = [v for v in vendas if v["Status"] == "Fechado Ganho"]
    assert len(vendas) == 29
    assert sum(v["Receita (R$)"] or 0 for v in vendas) == pytest.approx(123_810)
    assert sum(v["Receita (R$)"] or 0 for v in ganho) == pytest.approx(113_720)
    assert len(ganho) == 24


def test_arquivo_entregue_recalculado():
    """Confere os valores calculados no arquivo entregue (precisa ter sido aberto/salvo no Excel)."""
    if not ENTREGUE.exists():
        pytest.skip("planilha_vendas_tratada.xlsx ainda não foi gerada")
    rs = load_workbook(ENTREGUE, data_only=True)["Resumo"]
    if rs["B3"].value is None:
        pytest.skip("fórmulas sem valores calculados; abra e salve o arquivo no Excel")
    valores = [rs.cell(l, c).value for l in range(3, 20) for c in (2, 3)]
    assert not any(isinstance(x, str) and x.startswith("#") for x in valores), "erro de fórmula"
    assert rs["B3"].value == 29
    assert rs["B5"].value == pytest.approx(113_720)
    assert rs["B19"].value == pytest.approx(113_720)  # soma por região = receita ganha
